import os
import sys
import io
import openpyxl
from datetime import date, datetime
from django.core.files.uploadedfile import SimpleUploadedFile

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'school_management.settings')
import django
django.setup()

from django.test import Client, RequestFactory
from django.contrib.auth.models import AnonymousUser
from apps.accounts.models import User
from apps.academics.models import AcademicYear, Classroom
from apps.students.models import Student
from apps.teachers.models import Teacher
from apps.students.views import _calculate_age_grade_matrix

def run_tests():
    print("=" * 75)
    print("STARTING FULL MOEYS INTEGRATED VERIFICATION SUITE (ALL 4 AREAS)")
    print("=" * 75)

    # 1. Setup Admin Client
    admin_user = User.objects.filter(role='ADMIN', is_superuser=True).first()
    if not admin_user:
        admin_user = User.objects.filter(role='ADMIN').first()
    if not admin_user:
        admin_user = User.objects.create_superuser('admin_moeys_test', 'admin@test.com', 'Admin@123456', role='ADMIN')
    
    client = Client()
    client.force_login(admin_user)
    print(f"[Phase 1] Logged in as Admin: {admin_user.username} (Role: {admin_user.role})")

    # 2. Verify Academic Year 2026-2027 and Safe Fallbacks
    ay_2026 = AcademicYear.objects.filter(name='2026-2027').first()
    assert ay_2026 is not None, "Academic Year 2026-2027 must exist!"
    assert ay_2026.is_current is True, "Academic Year 2026-2027 must have is_current=True!"
    
    student_count_2026 = Student.objects.filter(academic_year=ay_2026).count()
    print(f"[Phase 2] Academic Year {ay_2026.name} (ID: {ay_2026.id}) verified: is_current={ay_2026.is_current}, enrolled students={student_count_2026}")
    assert student_count_2026 > 0, "Academic Year 2026-2027 must have enrolled students!"

    # 3. Test Student Age-Grade Statistics Matrix
    print("\n[Phase 3] Testing Student Age-Grade Statistics Calculation...")
    matrix_data = _calculate_age_grade_matrix(ay_2026)
    total_calc = matrix_data['total_students_count']
    total_female = matrix_data['total_female_count']
    pct_female = matrix_data['female_percent']
    print(f"  ✓ Total students in matrix: {total_calc} (Expected: {student_count_2026})")
    print(f"  ✓ Total female students: {total_female} ({pct_female}%)")
    assert total_calc == student_count_2026, f"Matrix total ({total_calc}) must match DB count ({student_count_2026})!"
    assert total_female > 0, "Female count must be > 0"

    # Test HTTP Web View for Age-Grade Statistics
    resp_stats = client.get('/students/statistics/age-grade/')
    assert resp_stats.status_code == 200, f"Age-Grade Statistics view failed with code {resp_stats.status_code}"
    assert 'ស្ថិតិសិស្សតាមអាយុ និងកម្រិតថ្នាក់' in resp_stats.content.decode('utf-8')
    print("  ✓ HTTP 200 OK: /students/statistics/age-grade/ loaded successfully")

    # 4. Test Web Upload & Sync of MoEYS Individual Student Extract (35 Columns)
    print("\n[Phase 4] Testing Web Upload & Sync for MoEYS Individual Student Census...")
    test_sid = "TEST999901"
    Student.objects.filter(student_id=test_sid).delete()

    # Build a valid 35-column MoEYS Census in memory
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "សរុប"
    # Row 6: First student
    ws.cell(6, 1, 1) # No
    ws.cell(6, 2, test_sid) # Student ID
    ws.cell(6, 3, "សុខ") # Surname
    ws.cell(6, 4, "ចរិយា") # Given name
    ws.cell(6, 5, "ស្រី") # Gender
    ws.cell(6, 6, 15) # Day
    ws.cell(6, 7, 8) # Month
    ws.cell(6, 8, 2009) # Year
    ws.cell(6, 9, "ដើមឫស") # Commune
    ws.cell(6, 10, "កណ្តាលស្ទឹង") # District
    ws.cell(6, 11, "កណ្តាល") # Province
    ws.cell(6, 12, 12) # Grade
    ws.cell(6, 13, "A") # Section
    ws.cell(6, 14, "សុខ វិបុល") # Father
    ws.cell(6, 15, "កសិករ") # Father Job
    ws.cell(6, 16, "ស៊ឹម សុភា") # Mother
    ws.cell(6, 17, "មេផ្ទះ") # Mother Job
    ws.cell(6, 20, "នៅជាមួយឪពុកម្តាយ") # Orphan
    ws.cell(6, 27, "ក្រ១") # Equity 1
    ws.cell(6, 31, "012345678") # Phone
    ws.cell(6, 32, "ACTIVE")
    ws.cell(6, 33, True) # Science Track
    
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)

    uploaded_file = SimpleUploadedFile(
        "test_moeys_individual_extract.xlsx",
        excel_stream.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    resp_upload = client.post('/students/reports/moeys-individual-roster/upload/', {
        'academic_year': str(ay_2026.id),
        'file': uploaded_file
    }, follow=True)

    assert resp_upload.status_code == 200, f"Upload failed with code {resp_upload.status_code}"
    print("  ✓ HTTP 200 OK after redirect: /students/reports/moeys-individual-roster/upload/")

    # Verify student was created in database with 35-column enrollment_data
    synced_st = Student.objects.filter(student_id=test_sid).first()
    assert synced_st is not None, f"Student {test_sid} must be created in database!"
    assert synced_st.khmer_name == "សុខ ចរិយា", f"Student name must be 'សុខ ចរិយា', got {synced_st.khmer_name}"
    assert synced_st.gender == 'F', f"Student gender must be 'F', got {synced_st.gender}"
    assert synced_st.date_of_birth == date(2009, 8, 15), f"DOB must match, got {synced_st.date_of_birth}"
    assert synced_st.enrollment_data.get('equity_card_1') == "ក្រ១", "Equity card must be 'ក្រ១'"
    assert synced_st.enrollment_data.get('track') == "វិទ្យាសាស្ត្រ", "Track must be 'វិទ្យាសាស្ត្រ'"
    print(f"  ✓ Verified Student {test_sid} created with 35-column MoEYS profile and science track")

    # Clean up test student
    synced_st.delete()

    # 5. Test Web Upload & Sync of MoEYS Staff Roster (2026.xlsx)
    print("\n[Phase 5] Testing Web Upload & Sync for MoEYS Staff Roster (2026.xlsx)...")
    resp_staff_page = client.get('/teachers/moeys-staff-roster/')
    assert resp_staff_page.status_code == 200, f"Staff roster view failed with {resp_staff_page.status_code}"
    staff_content = resp_staff_page.content.decode('utf-8')
    assert 'បញ្ជីគ្រប់គ្រងបុគ្គលិក មន្ត្រីរាជការ' in staff_content
    assert 'uploadMoEYSStaffModal' in staff_content, "Modal #uploadMoEYSStaffModal must be present in HTML!"
    print("  ✓ Staff Roster page loads with Modal #uploadMoEYSStaffModal")

    # Build a test staff file
    wb_staff = openpyxl.Workbook()
    ws_staff = wb_staff.active
    ws_staff.title = "2026-2027"
    test_tid = "TTEST9901"
    ws_staff.cell(8, 1, 1)
    ws_staff.cell(8, 2, test_tid)
    ws_staff.cell(8, 3, "ចាន់ វណ្ណា")
    ws_staff.cell(8, 4, "ស") # Female
    ws_staff.cell(8, 5, date(1985, 5, 20))
    ws_staff.cell(8, 6, "បរិញ្ញាបត្រ")
    ws_staff.cell(8, 7, "គណិតវិទ្យា")
    ws_staff.cell(8, 8, "គរុកោសល្យជាន់ខ្ពស់")
    ws_staff.cell(8, 11, "គណិតវិទ្យា")
    ws_staff.cell(8, 13, "គ្រូបង្រៀន")
    ws_staff.cell(8, 20, "098765432")

    staff_stream = io.BytesIO()
    wb_staff.save(staff_stream)
    staff_stream.seek(0)

    uploaded_staff = SimpleUploadedFile(
        "test_2026.xlsx",
        staff_stream.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    resp_staff_up = client.post('/teachers/moeys-staff-roster/upload/', {
        'file': uploaded_staff
    }, follow=True)
    assert resp_staff_up.status_code == 200, f"Staff upload failed with {resp_staff_up.status_code}"

    # Verify teacher record and user account were created
    synced_teacher = Teacher.objects.filter(teacher_id=test_tid).first()
    assert synced_teacher is not None, f"Teacher {test_tid} must be created!"
    assert synced_teacher.khmer_name == "ចាន់ វណ្ណា"
    assert synced_teacher.gender == 'F'
    assert synced_teacher.user is not None, "User login account must be created for teacher!"
    assert synced_teacher.user.username == test_tid
    print(f"  ✓ Verified Teacher {test_tid} created and synced with user login account")

    # Clean up test teacher & user
    synced_teacher.user.delete()
    synced_teacher.delete()

    # Restore root 2026.xlsx if needed from git
    os.system("git checkout -- 2026.xlsx")

    # 6. Test Excel Exports Across All 4 Views
    print("\n[Phase 6] Testing Excel (.xlsx) Downloads across all 4 MoEYS Views...")
    
    # 6a. MoEYS Individual Roster Excel
    r_exp1 = client.get('/students/reports/moeys-individual-roster/export-excel/')
    assert r_exp1.status_code == 200
    assert 'spreadsheetml.sheet' in r_exp1['Content-Type']
    print(f"  ✓ Export 1 (MoEYS Individual 35-Col .xlsx): HTTP 200, {len(r_exp1.content)} bytes")

    # 6b. MoEYS Staff Roster Excel
    r_exp2 = client.get('/teachers/moeys-staff-roster/export-excel/')
    assert r_exp2.status_code == 200
    assert 'spreadsheetml.sheet' in r_exp2['Content-Type']
    print(f"  ✓ Export 2 (MoEYS Staff Roster .xlsx): HTTP 200, {len(r_exp2.content)} bytes")

    # 6c. Student Age Custom Roster Excel
    r_exp3 = client.get('/students/reports/age-roster/export-excel/?format=format_a&mode=range&min_age=13&max_age=14')
    assert r_exp3.status_code == 200
    assert 'spreadsheetml.sheet' in r_exp3['Content-Type']
    print(f"  ✓ Export 3 (Student Age Custom Roster Format A .xlsx): HTTP 200, {len(r_exp3.content)} bytes")

    # 6d. Student Age-Grade Statistics Multi-Sheet Excel
    r_exp4 = client.get('/students/statistics/age-grade/export-excel/')
    assert r_exp4.status_code == 200
    assert 'spreadsheetml.sheet' in r_exp4['Content-Type']
    print(f"  ✓ Export 4 (MoEYS Age-Grade Matrix Multi-Sheet .xlsx): HTTP 200, {len(r_exp4.content)} bytes")

    # 7. Test Print / PDF View Rendering
    print("\n[Phase 7] Testing Print / PDF Views (Margin 1.2cm & Drag Signatures)...")
    
    # 7a. MoEYS Individual Roster Print
    r_prn1 = client.get('/students/reports/moeys-individual-roster/print/')
    assert r_prn1.status_code == 200
    c1 = r_prn1.content.decode('utf-8')
    assert '@page' in c1 and '1.2cm' in c1, "Print layout must contain @page margin 1.2cm!"
    assert 'signature-container' in c1, "Must contain draggable signature container!"
    print("  ✓ Print View 1 (MoEYS Individual Roster Print): HTTP 200, 1.2cm margin & draggable signatures verified")

    # 7b. MoEYS Staff Roster Print
    r_prn2 = client.get('/teachers/moeys-staff-roster/print/')
    assert r_prn2.status_code == 200
    c2 = r_prn2.content.decode('utf-8')
    assert '@page' in c2 and '1.2cm' in c2
    print("  ✓ Print View 2 (MoEYS Staff Roster Print): HTTP 200, 1.2cm margin verified")

    # 7c. Student Age Custom Roster Print
    r_prn3 = client.get('/students/reports/age-roster/print/?format=format_a&mode=range&min_age=13&max_age=14')
    assert r_prn3.status_code == 200
    c3 = r_prn3.content.decode('utf-8')
    assert '@page' in c3 and '1.2cm' in c3
    print("  ✓ Print View 3 (Student Age Custom Roster Print): HTTP 200, 1.2cm margin verified")

    print("\n" + "=" * 75)
    print("ALL 7 PHASES PASSED WITH 100% SUCCESS!")
    print("=" * 75)

if __name__ == '__main__':
    run_tests()
