r"""
Management command to synchronize student census records from
E:\SchoolSM\សម្រង់ព័ត៌មានសិស្សម្នាក់ៗ.xlsx into Student model and enrollment_data JSONField.
"""

import os
import openpyxl
from datetime import date
from django.core.management.base import BaseCommand
from apps.students.models import Student
from apps.academics.models import AcademicYear, Classroom

class Command(BaseCommand):
    help = "Synchronizes student records from E:\\SchoolSM\\សម្រង់ព័ត៌មានសិស្សម្នាក់ៗ.xlsx into the database"

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default=r'E:\SchoolSM\សម្រង់ព័ត៌មានសិស្សម្នាក់ៗ.xlsx',
            help='Path to the MoEYS individual student information Excel file'
        )

    def handle(self, *args, **options):
        file_path = options['file']
        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"File not found: {file_path}"))
            return

        self.stdout.write(self.style.NOTICE(f"Loading Excel file: {file_path}..."))
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        updated_count = 0
        created_count = 0
        skipped_count = 0

        # Retrieve active academic year 2026-2027
        ay = AcademicYear.objects.filter(name__icontains='2026-2027').exclude(name__icontains='Test').first()
        if not ay:
            ay = AcademicYear.objects.filter(is_current=True).first()

        for r in range(6, ws.max_row + 1):
            student_id_val = ws.cell(r, 2).value
            if not student_id_val:
                continue

            student_id = str(student_id_val).strip()
            last_name = str(ws.cell(r, 3).value or '').strip()
            first_name = str(ws.cell(r, 4).value or '').strip()
            full_khmer_name = f"{last_name} {first_name}".strip() if (last_name or first_name) else ''
            
            gender_val = str(ws.cell(r, 5).value or '').strip()
            gender_code = 'F' if ('ស្រី' in gender_val or gender_val.upper() == 'F') else 'M'

            # DOB
            dob_day = ws.cell(r, 6).value
            dob_month = ws.cell(r, 7).value
            dob_year = ws.cell(r, 8).value
            dob_obj = None
            if dob_day and dob_month and dob_year:
                try:
                    dob_obj = date(int(dob_year), int(dob_month), int(dob_day))
                except Exception:
                    pass

            # POB
            pob_commune = str(ws.cell(r, 9).value or '').strip()
            pob_district = str(ws.cell(r, 10).value or '').strip()
            pob_province = str(ws.cell(r, 11).value or '').strip()
            pob_parts = [p for p in [pob_commune, pob_district, pob_province] if p]
            place_of_birth = ", ".join(pob_parts) if pob_parts else None

            # Classroom
            grade_num = ws.cell(r, 12).value
            class_letter = str(ws.cell(r, 13).value or '').strip().upper()
            classroom = None
            if grade_num and class_letter:
                c_code = f"{grade_num}{class_letter}"
                classroom = Classroom.objects.filter(academic_year=ay, code__iexact=c_code).first()
                if not classroom:
                    classroom = Classroom.objects.filter(academic_year=ay, name__icontains=c_code).first()

            # Parents
            father_name = str(ws.cell(r, 14).value or '').strip() or None
            father_job = str(ws.cell(r, 15).value or '').strip() or None
            mother_name = str(ws.cell(r, 16).value or '').strip() or None
            mother_job = str(ws.cell(r, 17).value or '').strip() or None
            guardian_name = str(ws.cell(r, 18).value or '').strip() or None
            guardian_job = str(ws.cell(r, 19).value or '').strip() or None

            # Extended Profile Fields (stored in enrollment_data)
            orphan_status = str(ws.cell(r, 20).value or '').strip()
            primary_school = str(ws.cell(r, 21).value or '').strip()
            secondary_school = str(ws.cell(r, 22).value or '').strip()
            ethnic_minority = str(ws.cell(r, 23).value or '').strip()
            disability_physical = str(ws.cell(r, 24).value or '').strip()
            disability_sight = str(ws.cell(r, 25).value or '').strip()
            disability_hearing = str(ws.cell(r, 26).value or '').strip()
            equity_card_1 = str(ws.cell(r, 27).value or '').strip()
            equity_card_2 = str(ws.cell(r, 28).value or '').strip()
            risk_card = str(ws.cell(r, 29).value or '').strip()
            scholarship = str(ws.cell(r, 30).value or '').strip()
            
            # Phone
            phone_val = ws.cell(r, 31).value
            phone_str = ''
            if phone_val is not None:
                phone_raw = str(phone_val).strip()
                if phone_raw:
                    phone_str = f"0{phone_raw}" if not phone_raw.startswith('0') else phone_raw

            status_str = str(ws.cell(r, 32).value or '').strip() or 'ACTIVE'

            # Tracks
            is_sc = bool(ws.cell(r, 33).value)
            is_ss = bool(ws.cell(r, 34).value)
            is_voc = bool(ws.cell(r, 35).value)
            track_name = 'វិទ្យាសាស្ត្រ' if is_sc else ('វិទ្យាសាស្ត្រសង្គម' if is_ss else ('វិជ្ជាជីវៈ' if is_voc else 'ទូទៅ'))

            # Build enrollment_data dictionary
            moeys_profile = {
                'surname': last_name,
                'given_name': first_name,
                'pob_commune': pob_commune,
                'pob_district': pob_district,
                'pob_province': pob_province,
                'father_job': father_job,
                'mother_job': mother_job,
                'guardian_name': guardian_name,
                'guardian_job': guardian_job,
                'orphan_status': orphan_status,
                'primary_school': primary_school,
                'secondary_school': secondary_school,
                'ethnic_minority': ethnic_minority,
                'disability_physical': disability_physical,
                'disability_sight': disability_sight,
                'disability_hearing': disability_hearing,
                'equity_card_1': equity_card_1,
                'equity_card_2': equity_card_2,
                'risk_card': risk_card,
                'scholarship': scholarship,
                'track': track_name,
                'is_sc': is_sc,
                'is_ss': is_ss,
                'is_voc': is_voc,
            }

            # Find student by student_id or create
            student = Student.objects.filter(student_id=student_id).first()
            if student:
                # Update student fields
                if full_khmer_name and not student.khmer_name:
                    student.khmer_name = full_khmer_name
                if dob_obj and not student.date_of_birth:
                    student.date_of_birth = dob_obj
                if gender_code and not student.gender:
                    student.gender = gender_code
                if place_of_birth:
                    student.place_of_birth = place_of_birth
                if father_name:
                    student.father_name = father_name
                if father_job:
                    student.father_job = father_job
                if mother_name:
                    student.mother_name = mother_name
                if mother_job:
                    student.mother_job = mother_job
                if guardian_name:
                    student.guardian_name = guardian_name
                if phone_str:
                    student.phone = phone_str
                if classroom and not student.classroom:
                    student.classroom = classroom
                if ay and not student.academic_year:
                    student.academic_year = ay

                # Merge enrollment_data
                existing_ed = dict(student.enrollment_data or {})
                existing_ed.update(moeys_profile)
                student.enrollment_data = existing_ed
                student.save()
                updated_count += 1
            else:
                # Create student
                ed = dict(moeys_profile)
                Student.objects.create(
                    student_id=student_id,
                    khmer_name=full_khmer_name or f"សិស្ស {student_id}",
                    gender=gender_code,
                    date_of_birth=dob_obj or date(2008, 1, 1),
                    place_of_birth=place_of_birth,
                    classroom=classroom,
                    academic_year=ay,
                    father_name=father_name,
                    father_job=father_job,
                    mother_name=mother_name,
                    mother_job=mother_job,
                    guardian_name=guardian_name,
                    phone=phone_str or None,
                    status='ACTIVE',
                    enrollment_data=ed
                )
                created_count += 1

        self.stdout.write(self.style.SUCCESS(
            f"Sync complete: {updated_count} students updated, {created_count} students created, {skipped_count} skipped."
        ))
