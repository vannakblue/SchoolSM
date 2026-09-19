import os
import sys
import io
import django
import openpyxl

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'school_management.settings')
django.setup()

from django.test import Client
from apps.accounts.models import User
from apps.academics.models import AcademicYear, Classroom

def run_tests():
    print("==========================================================================")
    print("🚀 TEST: ADMIN EXCEL EXPORT FOR STUDENT & TEACHER ATTENDANCE REPORTS")
    print("==========================================================================")

    admin = User.objects.filter(role='ADMIN').first()
    if not admin:
        admin = User.objects.create_superuser('admin_test_export', 'admin@export.com', 'pass123')

    teacher = User.objects.filter(role='TEACHER').first()
    if not teacher:
        teacher = User.objects.create_user('teacher_test_export', 'teacher@export.com', 'pass123', role='TEACHER')

    c_admin = Client()
    c_admin.force_login(admin)

    c_teacher = Client()
    c_teacher.force_login(teacher)

    # --------------------------------------------------------------------------
    # 1. Student Attendance Report Excel Export
    # --------------------------------------------------------------------------
    print("\n--- 1. Testing Student Attendance Report Excel Export ---")
    resp_stu_all = c_admin.get('/attendance/report/?export=excel')
    assert resp_stu_all.status_code == 200, f"Expected 200, got {resp_stu_all.status_code}"
    assert 'spreadsheetml.sheet' in resp_stu_all.get('Content-Type')
    
    wb_stu = openpyxl.load_workbook(io.BytesIO(resp_stu_all.content))
    ws_stu = wb_stu.active
    assert ws_stu.title == "វត្តមានសិស្ស"
    assert "STUDENT ATTENDANCE REPORT" in ws_stu['A2'].value
    assert "សិស្សសរុប" in ws_stu['A4'].value
    expected_stu_headers = ['ល.រ', 'អត្តលេខសិស្ស', 'គោត្តនាម និងនាម', 'ឈ្មោះឡាតាំង', 'ថ្នាក់រៀន', 'ភេទ', 'វត្តមាន (Present)', 'ច្បាប់ (Permission)', 'អវត្តមាន (Absent)', 'មកយឺត (Late)', 'វេនសរុប (Total)', 'អត្រាវត្តមាន (%)']
    actual_stu_headers = [ws_stu.cell(row=6, column=i).value for i in range(1, 13)]
    assert actual_stu_headers == expected_stu_headers, f"Header mismatch: {actual_stu_headers}"
    print(f"✅ Student Attendance Report (All Classes) Exported: Rows={ws_stu.max_row}, Columns={ws_stu.max_column}")

    # Test with specific classroom filter
    target_class = Classroom.objects.first()
    if target_class:
        resp_stu_cls = c_admin.get(f'/attendance/report/?classroom={target_class.id}&filter_type=month&export=excel')
        assert resp_stu_cls.status_code == 200
        wb_cls = openpyxl.load_workbook(io.BytesIO(resp_stu_cls.content))
        ws_cls = wb_cls.active
        assert "ថ្នាក់រៀន៖" in ws_cls['A3'].value or "វិសាលភាព៖" in ws_cls['A3'].value
        print(f"✅ Student Attendance Report (Classroom {target_class.name}) Exported Successfully")

    # --------------------------------------------------------------------------
    # 2. Teacher Attendance Report Excel Export
    # --------------------------------------------------------------------------
    print("\n--- 2. Testing Teacher Attendance Report Excel Export ---")
    # A. Month Mode
    resp_tch_month = c_admin.get('/teachers/attendance/report/?filter_type=month&export=excel')
    assert resp_tch_month.status_code == 200
    wb_tch_m = openpyxl.load_workbook(io.BytesIO(resp_tch_month.content))
    ws_tch_m = wb_tch_m.active
    assert ws_tch_m.title == "វត្តមានគ្រូបង្រៀន"
    assert "TEACHER ATTENDANCE REPORT" in ws_tch_m['A2'].value
    assert "គ្រូសរុប" in ws_tch_m['A4'].value
    expected_tch_m_headers = ['ល.រ', 'អត្តលេខ', 'គោត្តនាម និងនាម', 'ឈ្មោះឡាតាំង', 'ភេទ', 'ឯកទេស', 'ថ្ងៃត្រូវបង្រៀន', 'ម៉ោងត្រូវបង្រៀន', 'ម៉ោងបានស្រង់', 'ម៉ោងខកខាន', 'ថ្ងៃសុំច្បាប់', 'អត្រាអនុលោមភាព (%)']
    actual_tch_m_headers = [ws_tch_m.cell(row=6, column=i).value for i in range(1, 13)]
    assert actual_tch_m_headers == expected_tch_m_headers, f"Header mismatch: {actual_tch_m_headers}"
    print(f"✅ Teacher Attendance Report (Month Mode) Exported: Rows={ws_tch_m.max_row}, Columns={ws_tch_m.max_column}")

    # B. Daily Mode
    resp_tch_day = c_admin.get('/teachers/attendance/report/?filter_type=day&export=excel')
    assert resp_tch_day.status_code == 200
    wb_tch_d = openpyxl.load_workbook(io.BytesIO(resp_tch_day.content))
    ws_tch_d = wb_tch_d.active
    assert "TEACHER ATTENDANCE REPORT" in ws_tch_d['A2'].value
    assert "គ្រូត្រូវបង្រៀនសរុប" in ws_tch_d['A4'].value
    expected_tch_d_headers = ['ល.រ', 'អត្តលេខ', 'គោត្តនាម និងនាម', 'ឈ្មោះឡាតាំង', 'ភេទ', 'ឯកទេស', 'ម៉ោងទី ១', 'ម៉ោងទី ២', 'ម៉ោងទី ៣', 'ម៉ោងទី ៤', 'ម៉ោងទី ៥', 'ម៉ោងទី ៦', 'ម៉ោងទី ៧', 'ម៉ោងទី ៨', 'ម៉ោងត្រូវបង្រៀន', 'បានស្រង់', 'ខកខាន', 'ស្ថានភាព']
    actual_tch_d_headers = [ws_tch_d.cell(row=6, column=i).value for i in range(1, 19)]
    assert actual_tch_d_headers == expected_tch_d_headers, f"Header mismatch: {actual_tch_d_headers}"
    print(f"✅ Teacher Attendance Report (Daily Mode with P1-P8) Exported: Rows={ws_tch_d.max_row}, Columns={ws_tch_d.max_column}")

    # --------------------------------------------------------------------------
    # 3. Security & Role Restriction Verification
    # --------------------------------------------------------------------------
    print("\n--- 3. Testing Permission Control (Admin Only) ---")
    resp_non_admin_stu = c_teacher.get('/attendance/report/?export=excel')
    assert resp_non_admin_stu.status_code == 403, f"Expected 403 Forbidden for teacher, got {resp_non_admin_stu.status_code}"

    resp_non_admin_tch = c_teacher.get('/teachers/attendance/report/?export=excel')
    assert resp_non_admin_tch.status_code == 403, f"Expected 403 Forbidden for teacher, got {resp_non_admin_tch.status_code}"
    print("✅ Non-admin (Teacher) correctly blocked with 403 Forbidden on both reports.")

    print("\n==========================================================================")
    print("🎉 ALL ADMIN ATTENDANCE EXCEL EXPORT TESTS PASSED 100%!")
    print("==========================================================================")

if __name__ == '__main__':
    run_tests()
