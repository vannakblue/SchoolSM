import os
import io
import sys
import django
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'school_management.settings')
django.setup()

from django.test.utils import setup_test_environment
setup_test_environment()

from django.test import Client
from django.urls import reverse
from apps.accounts.models import User

def test_moeys_staff_roster():
    print("================================================================================")
    print("🧪 TESTING MOEYS CIVIL SERVANT & TEACHER DIRECTORY 2026-2027 (2026.xlsx)")
    print("================================================================================")

    # 1. Admin login
    admin_user = User.objects.filter(role='ADMIN').first()
    if not admin_user:
        admin_user = User.objects.create_superuser('admin_roster_test', 'admin_roster@test.com', 'adminpass123')

    client = Client()
    client.force_login(admin_user)

    # 2. Test Main View Rendering (Status 200)
    print("\n--- PHASE 1: Main View Rendering & Data Completeness ---")
    url = reverse('moeys_staff_roster')
    resp = client.get(url)
    assert resp.status_code == 200, f"Failed to load roster page: {resp.status_code}"
    content = resp.content.decode('utf-8')
    assert "បញ្ជីគ្រប់គ្រងបុគ្គលិក មន្ត្រីរាជការ ឆ្នាំសិក្សា ២០២៦-២០២៧" in content
    assert "វិទ្យាល័យ ហ៊ុន សែន កំពង់កន្ទួត" in content
    assert "ផេង រិទ្ធីយ៉ា" in content
    assert "ទិត សោម៉នវីរៈ" in content
    assert "ទឹម ប៊ុនធន" in content
    assert "1720800778" in content
    # Total count in context
    staff_count = len(resp.context['staff_list'])
    print(f"Total staff members rendered: {staff_count}")
    assert staff_count == 119, f"Expected 119 staff members, got {staff_count}"
    assert resp.context['total_female'] == 64, f"Expected 64 female staff, got {resp.context['total_female']}"
    assert resp.context['total_male'] == 55, f"Expected 55 male staff, got {resp.context['total_male']}"
    assert resp.context['count_cat_a'] == 60, f"Expected 60 Category A, got {resp.context['count_cat_a']}"
    assert resp.context['count_cat_b'] == 54, f"Expected 54 Category B, got {resp.context['count_cat_b']}"
    assert resp.context['count_cat_c'] == 5, f"Expected 5 Category C, got {resp.context['count_cat_c']}"
    print("✅ [PASS] Main Web Roster rendered 119 staff members with exact framework distribution!")

    # 3. Test Search & Filters
    print("\n--- PHASE 2: Live Search & Framework Filters ---")
    # Filter by Category A (ក) -> 60
    resp_a = client.get(f"{url}?cat=ក")
    assert resp_a.status_code == 200
    assert len(resp_a.context['staff_list']) == 60, f"Expected 60 for Category A, got {len(resp_a.context['staff_list'])}"
    print(f"Filter [Category ក]: {len(resp_a.context['staff_list'])} members matched.")

    # Filter by Category B (ខ) -> 54
    resp_b = client.get(f"{url}?cat=ខ")
    assert resp_b.status_code == 200
    assert len(resp_b.context['staff_list']) == 54, f"Expected 54 for Category B, got {len(resp_b.context['staff_list'])}"
    print(f"Filter [Category ខ]: {len(resp_b.context['staff_list'])} members matched.")

    # Filter by Category C (គ) -> 5
    resp_c = client.get(f"{url}?cat=គ")
    assert resp_c.status_code == 200
    assert len(resp_c.context['staff_list']) == 5, f"Expected 5 for Category C, got {len(resp_c.context['staff_list'])}"
    print(f"Filter [Category គ]: {len(resp_c.context['staff_list'])} members matched.")

    # Filter by Gender Female (ស) -> 64
    resp_fem = client.get(f"{url}?gender=ស")
    assert resp_fem.status_code == 200
    assert len(resp_fem.context['staff_list']) == 64, f"Expected 64 female, got {len(resp_fem.context['staff_list'])}"
    print(f"Filter [Female]: {len(resp_fem.context['staff_list'])} female staff matched.")

    # Search Query
    resp_q = client.get(f"{url}?q=ផេង រិទ្ធីយ៉ា")
    assert resp_q.status_code == 200
    assert len(resp_q.context['staff_list']) == 1
    assert resp_q.context['staff_list'][0]['name'] == 'ផេង រិទ្ធីយ៉ា'
    print("Filter [Search ផេង រិទ្ធីយ៉ា]: Principal found successfully.")
    print("✅ [PASS] All search and filtering options verified!")

    # 4. Test Excel Download
    print("\n--- PHASE 3: Excel (.xlsx) Export Capability ---")
    url_excel = reverse('moeys_staff_roster_export_excel')
    resp_excel = client.get(url_excel)
    assert resp_excel.status_code == 200, f"Failed to download Excel: {resp_excel.status_code}"
    assert resp_excel['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment' in resp_excel['Content-Disposition']
    assert 'moeys_staff_roster_2026_2027.xlsx' in resp_excel['Content-Disposition']

    wb = openpyxl.load_workbook(io.BytesIO(resp_excel.content))
    assert '2026-2027' in wb.sheetnames
    ws = wb['2026-2027']
    print(f"Downloaded Excel Sheet: {ws.title}, Max Row: {ws.max_row}, Max Col: {ws.max_column}")
    assert ws.max_row >= 126, "Must contain all 119 teachers"
    assert ws['C8'].value == 'ផេង រិទ្ធីយ៉ា'
    assert ws['B8'].value == 1720800778
    print("✅ [PASS] Excel file downloaded successfully matching 2026.xlsx!")

    # 5. Test Print / PDF Saving View
    print("\n--- PHASE 4: Dedicated Print & Save as PDF View ---")
    url_print = reverse('moeys_staff_roster_print')
    resp_print = client.get(url_print)
    assert resp_print.status_code == 200, f"Failed to load print view: {resp_print.status_code}"
    content_print = resp_print.content.decode('utf-8')
    assert "បញ្ជីគ្រប់គ្រងបុគ្គលិក មន្ត្រីរាជការ ឆ្នាំសិក្សា ២០២៦-២០២៧" in content_print
    assert "@page" in content_print
    assert "A3 landscape" in content_print
    assert "window.print()" in content_print
    assert "ផេង រិទ្ធីយ៉ា" in content_print
    assert "ប្រធានការិយាល័យអប់រំ យុវជន និងកីឡាស្រុក" in content_print
    print("✅ [PASS] Print view rendered with A3 landscape print CSS and official MoEYS footer!")

    # 6. Test Teacher Role Access
    print("\n--- PHASE 5: Role & Permission Verification ---")
    teacher_user = User.objects.filter(role='TEACHER').first()
    if teacher_user:
        client.force_login(teacher_user)
        resp_tch = client.get(url)
        assert resp_tch.status_code == 200, "Teachers must have access to view the staff directory"
        print(f"✅ [PASS] Teacher [{teacher_user.username}] successfully accessed the Staff Roster!")

    print("\n================================================================================")
    print("🎉 ALL 5 TEST PHASES PASSED WITH 100% SUCCESS!")
    print("================================================================================")

if __name__ == '__main__':
    test_moeys_staff_roster()
