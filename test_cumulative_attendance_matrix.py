import os
import sys
import django
import io
import openpyxl

if sys.stdout and sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'school_management.settings')
django.setup()

from django.test import RequestFactory
from django.contrib.auth import get_user_model
from apps.attendance.views import student_semester_annual_attendance_report
from apps.accounts.models import User

def run_tests():
    print("=== Testing Cumulative Semester/Annual Attendance Matrix & Excel Export ===")
    factory = RequestFactory()

    # Find or get admin and non-admin users
    admin_user = User.objects.filter(role='ADMIN').first()
    if not admin_user:
        admin_user = User.objects.filter(is_superuser=True).first()
    if not admin_user:
        admin_user = User.objects.create_superuser('testadmin', 'admin@example.com', 'pass123', role='ADMIN')

    teacher_user = User.objects.filter(role='TEACHER').first()
    if not teacher_user:
        teacher_user = User.objects.filter(role__in=['TEACHER', 'STAFF']).first()
    if not teacher_user:
        teacher_user = User.objects.create_user('testteacher', 'teacher@example.com', 'pass123', role='TEACHER')

    print(f"Admin User: {admin_user.username} ({admin_user.role})")
    print(f"Teacher User: {teacher_user.username} ({teacher_user.role})")

    # 1. Test HTML View: CLASS_SUMMARY (Default)
    req1 = factory.get('/attendance/semester-annual-summary/')
    req1.user = admin_user
    resp1 = student_semester_annual_attendance_report(req1)
    assert resp1.status_code == 200, f"Expected 200, got {resp1.status_code}"
    content1 = resp1.content.decode('utf-8')
    assert "របាយការណ៍អវត្តមានសិស្ស តាមខែ/ឆមាស/ប្រចាំឆ្នាំ" in content1, "Title missing in CLASS_SUMMARY"
    assert "ឆមាសទី ១ (SEMESTER 1)" in content1, "Semester 1 header missing"
    assert "ឆមាសទី ២ (SEMESTER 2)" in content1, "Semester 2 header missing"
    assert "សរុបប្រចាំឆ្នាំ" in content1, "Annual header missing"
    print("✓ Test 1 Passed: CLASS_SUMMARY HTML view returns 200 with complete matrix headers")

    # 2. Test HTML View: STUDENT_ROSTER
    req2 = factory.get('/attendance/semester-annual-summary/?scope=STUDENT_ROSTER&unit=DAYS&absence_type=DUAL')
    req2.user = admin_user
    resp2 = student_semester_annual_attendance_report(req2)
    assert resp2.status_code == 200, f"Expected 200, got {resp2.status_code}"
    content2 = resp2.content.decode('utf-8')
    assert "គោត្តនាម-នាម" in content2, "Student roster header missing"
    assert "សរុបរួមសិស្សទាំងអស់" in content2 or "ពុំមានទិន្នន័យ" in content2, "Roster footer missing"
    print("✓ Test 2 Passed: STUDENT_ROSTER HTML view returns 200 with correct columns")

    # 3. Test Excel Export as Admin (CLASS_SUMMARY)
    req3 = factory.get('/attendance/semester-annual-summary/?scope=CLASS_SUMMARY&export=excel')
    req3.user = admin_user
    resp3 = student_semester_annual_attendance_report(req3)
    assert resp3.status_code == 200, f"Expected 200, got {resp3.status_code}"
    assert 'spreadsheetml' in resp3['Content-Type'], f"Invalid Content-Type: {resp3['Content-Type']}"
    assert 'attachment; filename=' in resp3['Content-Disposition'], "Missing attachment filename"
    
    wb3 = openpyxl.load_workbook(io.BytesIO(resp3.content))
    assert 'អវត្តមានខែ-ឆមាស-ប្រចាំឆ្នាំ' in wb3.sheetnames, f"Expected sheet not found: {wb3.sheetnames}"
    ws3 = wb3['អវត្តមានខែ-ឆមាស-ប្រចាំឆ្នាំ']
    print(f"✓ Test 3 Passed: Admin Excel export for CLASS_SUMMARY valid ({ws3.max_row} rows, {ws3.max_column} cols)")

    # 4. Test Excel Export as Admin (STUDENT_ROSTER)
    req4 = factory.get('/attendance/semester-annual-summary/?scope=STUDENT_ROSTER&export=excel')
    req4.user = admin_user
    resp4 = student_semester_annual_attendance_report(req4)
    assert resp4.status_code == 200, f"Expected 200, got {resp4.status_code}"
    wb4 = openpyxl.load_workbook(io.BytesIO(resp4.content))
    assert 'អវត្តមានខែ-ឆមាស-ប្រចាំឆ្នាំ' in wb4.sheetnames, f"Expected sheet not found: {wb4.sheetnames}"
    ws4 = wb4['អវត្តមានខែ-ឆមាស-ប្រចាំឆ្នាំ']
    print(f"✓ Test 4 Passed: Admin Excel export for STUDENT_ROSTER valid ({ws4.max_row} rows, {ws4.max_column} cols)")

    # 5. Security: Test Non-Admin Excel Export (Must be 403 Forbidden)
    req5 = factory.get('/attendance/semester-annual-summary/?export=excel')
    req5.user = teacher_user
    resp5 = student_semester_annual_attendance_report(req5)
    assert resp5.status_code == 403, f"Expected 403 for non-admin Excel export, got {resp5.status_code}"
    print("✓ Test 5 Passed: Non-admin users are strictly forbidden (403) from Excel export")

    # 6. Verification of Cumulative Math Integrity:
    # Test that get_academic_year_months_catalog correctly groups 9..2 as Sem 1 and 3..7 as Sem 2
    from apps.attendance.views import get_academic_year_months_catalog
    from apps.academics.models import AcademicYear
    ay = AcademicYear.objects.filter(is_current=True).first()
    cat = get_academic_year_months_catalog(ay)
    all_m = cat['all_months']
    s1_m = cat['sem1_months']
    s2_m = cat['sem2_months']
    assert len(all_m) == len(s1_m) + len(s2_m), "Total months count must equal sem1 + sem2"
    for m in s1_m:
        assert m['semester'] == 1
        assert m['month'] in [9, 10, 11, 12, 1, 2]
    for m in s2_m:
        assert m['semester'] == 2
        assert m['month'] in [3, 4, 5, 6, 7]
    print(f"✓ Test 6 Passed: Academic year months catalog verified (Sem1: {len(s1_m)} months, Sem2: {len(s2_m)} months, Total: {len(all_m)} months)")

    print("\nALL CUMULATIVE MATRIX & EXCEL TESTS PASSED SUCCESSFULLY! 🚀")

if __name__ == '__main__':
    run_tests()
